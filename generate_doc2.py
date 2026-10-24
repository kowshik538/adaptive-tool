import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUTPUT_DIR = r"C:\Users\user\.gemini\antigravity\scratch\adaptive-tool-scoping"
DOC2_PATH = os.path.join(OUTPUT_DIR, "Document_2_Page_by_Page_UI_and_Code_Implementation_Guide.docx")

COLOR_NAVY = RGBColor(12, 63, 110)
COLOR_SKY = RGBColor(2, 132, 199)
COLOR_SLATE = RGBColor(51, 65, 85)
COLOR_MUTED = RGBColor(100, 116, 139)
COLOR_DANGER = RGBColor(220, 38, 38)
COLOR_SUCCESS = RGBColor(22, 163, 74)

HEX_LIGHT_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_NAVY = "0C3F6E"
HEX_SKY_BG = "F0F9FF"
HEX_CODE_BG = "F8FAFC"
HEX_CALLOUT_BORDER = "0284C7"


def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_custom_heading(doc, text, level):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    p.paragraph_format.keep_with_next = True

    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = COLOR_NAVY
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(5)
    elif level == 2:
        run.font.size = Pt(12.5)
        run.font.color.rgb = COLOR_SKY
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_SLATE
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
    return p


def add_p(doc, text, bold_prefix="", space_after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.size = Pt(9.5)
        r_pre.font.color.rgb = COLOR_SLATE
    r_body = p.add_run(text)
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = COLOR_SLATE
    return p


def add_code_block(doc, code_str, title=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, HEX_CODE_BG)
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="0284C7"/>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05

    if title:
        t_run = p.add_run(f"💻 {title}\n")
        t_run.font.name = "Consolas"
        t_run.font.size = Pt(8.5)
        t_run.font.color.rgb = COLOR_NAVY
        t_run.bold = True

    c_run = p.add_run(code_str)
    c_run.font.name = "Consolas"
    c_run.font.size = Pt(8.0)
    c_run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_styled_table(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], HEX_NAVY)
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "FFFFFF" if r_idx % 2 == 0 else HEX_LIGHT_BG
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(8.0)
                run.font.color.rgb = COLOR_SLATE

    if col_widths:
        for row in table.rows:
            for c_idx, w in enumerate(col_widths):
                row.cells[c_idx].width = Inches(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(5)


def build_document_2():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title Block
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(16)
    p_title.paragraph_format.space_after = Pt(3)
    run_main_title = p_title.add_run("Adaptive Least-Privilege Tool Scoping: Complete Page-by-Page & Code Implementation Guide\n")
    run_main_title.bold = True
    run_main_title.font.size = Pt(20)
    run_main_title.font.color.rgb = COLOR_NAVY

    run_sub_title = p_title.add_run("Detailed Walkthrough of UI Pages, Interaction Workflows, Frontend React Code & Backend Python Engines\n")
    run_sub_title.bold = True
    run_sub_title.font.size = Pt(12)
    run_sub_title.font.color.rgb = COLOR_SKY

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(20)
    run_meta = p_meta.add_run("Document 2 of 2: Technical Architecture, UI Components & Source Code Reference")
    run_meta.font.size = Pt(9.5)
    run_meta.font.color.rgb = COLOR_MUTED

    # 1. Introduction
    add_custom_heading(doc, "1. System Architectural Organization", level=1)
    add_p(
        doc,
        "The project is structured as a full-stack, research-grade software platform consisting of a modular React Single-Page Application (SPA) on the frontend and an asynchronous Python FastAPI service on the backend, persisting data into a structured SQLite relational database. "
        "The web application delivers nine distinct functional pages and components, each serving a specific research or operational purpose as mandated by the Project Developer Brief."
    )

    add_styled_table(
        doc,
        ["Page / View", "Primary Objective", "Frontend File", "Backend Endpoints & Engine"],
        [
            ["Top Navigation & Role Switcher", "System branding, tab switching, and 1-click active test role impersonation", "src/components/Navbar.jsx", "/api/auth/me, auth.py"],
            ["1. Main Dashboard", "Executive research overview, KPI metrics, strategy comparison, and pipeline diagram", "src/pages/DashboardPage.jsx", "/api/metrics/summary, evaluation_engine.py"],
            ["2. Task Playground", "Interactive natural language task input, live 6-step scoping visualizer, and approval modal", "src/pages/PlaygroundPage.jsx", "/api/tasks/analyze, /api/tasks/execute, task_analyzer.py"],
            ["3. Tool Registry & Policies", "Full catalogue of 10 tools, risk scoring sliders, permitted verbs, and direct sandbox testing", "src/pages/ToolRegistryPage.jsx", "/api/tools, /api/tools/{name}/test, tool_registry.py"],
            ["4. Experiment Lab", "Controlled batch benchmark executions across Full/Moderate/Least-Privilege and 3B/70B models", "src/pages/ExperimentLabPage.jsx", "/api/experiments/run, benchmarks.py"],
            ["5. Security Testing Suite", "Controlled adversarial attack scenarios with side-by-side comparative exploit evaluations", "src/pages/SecurityTestingPage.jsx", "/api/security/test, agent_runner.py"],
            ["6. Tamper-Evident Audit Logs", "Filterable chronological audit journal with secret scrubbing and inspection drawer", "src/pages/AuditLogsPage.jsx", "/api/audit/logs, audit_logger.py"],
            ["7. Results & Reports", "Visual analytics for all 7 evaluation metrics with Recharts and CSV/JSON data download", "src/pages/ReportsPage.jsx", "/api/reports/export, evaluation_engine.py"],
            ["8. System Settings & RBAC", "Configurable risk thresholds, security toggles, and RBAC permission matrix", "src/pages/SettingsPage.jsx", "/api/settings, policy_engine.py"]
        ],
        [1.6, 2.2, 1.3, 1.4]
    )

    # 2. Navigation Bar & Role Switcher
    add_custom_heading(doc, "2. Navigation Bar & 1-Click Role Switcher", level=1)
    add_p(
        doc,
        "The Navigation Bar provides persistent header access across all views. It features an integrated 1-Click Role Switcher allowing examiners and developers to effortlessly toggle between Admin, Researcher, Tester, and Viewer roles without manual re-login, instantly demonstrating the RBAC boundary differences."
    )
    add_code_block(
        doc,
        "// src/components/Navbar.jsx - Role Switcher Dropdown Logic\n"
        "const roles = [\n"
        "  { name: 'Admin', user: 'admin', desc: 'Full permissions: edit tools, approve actions, modify DB' },\n"
        "  { name: 'Researcher', user: 'researcher', desc: 'Execute benchmarks, view analytics, query database' },\n"
        "  { name: 'Tester', user: 'tester', desc: 'Submit test tasks and evaluate security denials' },\n"
        "  { name: 'Viewer', user: 'viewer', desc: 'Read-only access to dashboards, reports, and logs' }\n"
        "];\n\n"
        "<button onClick={() => setCurrentUser({ ...currentUser, role: r.name, username: r.user })}>\n"
        "  <span className={r.color}>{r.name}</span>\n"
        "  <p className='text-[10px] text-slate-400'>{r.desc}</p>\n"
        "</button>",
        "Frontend: Navbar Role Switcher"
    )
    add_code_block(
        doc,
        "# backend/app/auth.py - Role-Based Access Control Dependency\n"
        "def require_role(min_role: UserRole):\n"
        "    hierarchy = { UserRole.VIEWER: 1, UserRole.TESTER: 2, UserRole.RESEARCHER: 3, UserRole.ADMIN: 4 }\n"
        "    def role_checker(user: dict = Depends(get_current_user_required)):\n"
        "        user_role = UserRole(user.get('role', UserRole.VIEWER.value))\n"
        "        if hierarchy.get(user_role, 1) < hierarchy.get(min_role, 1):\n"
        "            raise HTTPException(\n"
        "                status_code=status.HTTP_403_FORBIDDEN,\n"
        "                detail=f\"Operation requires '{min_role.value}' role privileges.\"\n"
        "            )\n"
        "        return user\n"
        "    return role_checker",
        "Backend: RBAC Dependency Decorator"
    )

    # 3. Main Dashboard Page
    add_custom_heading(doc, "3. Page 1: Main Dashboard (System Overview & KPIs)", level=1)
    add_p(
        doc,
        "The Main Dashboard summarizes research findings, displays top-level Key Performance Indicators (KPIs), highlights the comparative findings between the three access strategies, and visualizes the 6-stage enforcement pipeline."
    )
    add_code_block(
        doc,
        "// src/pages/DashboardPage.jsx - Loading Metrics and Calculating Exposure Reduction\n"
        "useEffect(() => {\n"
        "  async function loadData() {\n"
        "    const [sum, astats] = await Promise.all([api.getMetricsSummary(), api.getAuditStats()]);\n"
        "    setMetrics(sum);\n"
        "    setAuditStats(astats);\n"
        "  }\n"
        "  loadData();\n"
        "}, []);\n\n"
        "// Compute Exposure Reduction % between Full Access and Least-Privilege\n"
        "const fullExp = strategies['Full Access']?.avg_security_exposure || 45.0;\n"
        "const lpExp = strategies['Least-Privilege']?.avg_security_exposure || 3.5;\n"
        "const exposureReduction = Math.round(((fullExp - lpExp) / (fullExp || 1)) * 100);",
        "Frontend: Dashboard Metrics Calculation"
    )
    add_code_block(
        doc,
        "# backend/app/evaluation_engine.py - Aggregating Execution Logs into KPIs\n"
        "@classmethod\n"
        "def calculate_summary(cls, db: Session) -> Dict[str, Any]:\n"
        "    tasks = db.query(TaskExecutionModel).all()\n"
        "    df = pd.DataFrame([{\n"
        "        'strategy': t.access_strategy.value, 'model': t.model_scale.value,\n"
        "        'task_succeeded': 1 if t.task_succeeded else 0,\n"
        "        'attack_succeeded': 1 if t.attack_succeeded else 0,\n"
        "        'is_adversarial': 1 if t.is_adversarial_attempt else 0,\n"
        "        'security_exposure_score': t.security_exposure_score,\n"
        "        'latency_ms': t.latency_ms, 'total_tokens': t.prompt_tokens + t.completion_tokens\n"
        "    } for t in tasks])\n"
        "    # Attack Prevention Rate Calculation\n"
        "    adv_runs = df[df['is_adversarial'] == 1]\n"
        "    blocked = len(adv_runs[adv_runs['attack_succeeded'] == 0])\n"
        "    prevention_rate = round((blocked / len(adv_runs)) * 100, 1) if not adv_runs.empty else 100.0\n"
        "    return { 'attack_prevention_rate': prevention_rate, ... }",
        "Backend: Evaluation Engine Statistical Summary"
    )

    # 4. Task Playground
    add_custom_heading(doc, "4. Page 2: Task Playground & Live Scoping Inspector", level=1)
    add_p(
        doc,
        "The Task Playground allows users to test arbitrary task prompts or select preset benchmarks. Users choose the Access Strategy (Full, Moderate, Least-Privilege) and Model Scale (3B vs. 70B), then inspect the step-by-step enforcement pipeline in real time. It includes an interactive modal for human approval if a tool triggers REQUIRE_APPROVAL."
    )
    add_code_block(
        doc,
        "// src/pages/PlaygroundPage.jsx - Mediated Execution & Approval Handler\n"
        "const handleExecute = async () => {\n"
        "  setExecuting(true);\n"
        "  const res = await api.executeTask({\n"
        "    task_prompt: prompt,\n"
        "    access_strategy: strategy,\n"
        "    model_scale: modelScale,\n"
        "    override_role: roleOverride\n"
        "  });\n"
        "  setExecutionResult(res);\n"
        "  if (res.status === 'PENDING_APPROVAL' && res.approval_required_for) {\n"
        "    setApprovalModal({ taskId: res.task_id, toolName: res.approval_required_for });\n"
        "  }\n"
        "  setExecuting(false);\n"
        "};",
        "Frontend: Task Execution & Live Approval Modal Trigger"
    )
    add_code_block(
        doc,
        "# backend/app/scoping_engine.py - Dynamic Least-Privilege Allow-List Formulation\n"
        "@classmethod\n"
        "def compute_scope(cls, strategy: AccessStrategy, task_analysis: TaskAnalysisResult, user_role: UserRole):\n"
        "    scoped_allow_list = []\n"
        "    blocked_tools = []\n"
        "    if strategy == AccessStrategy.LEAST_PRIVILEGE:\n"
        "        for tool_name in task_analysis.candidate_tools:\n"
        "            eval_res = PolicyEngine.evaluate(tool_name, user_role, task_analysis, enforce_relevance=True)\n"
        "            if eval_res.decision in [PolicyDecision.ALLOW, PolicyDecision.REQUIRE_APPROVAL]:\n"
        "                scoped_allow_list.append(tool_name)\n"
        "            else:\n"
        "                blocked_tools.append(tool_name)\n"
        "    return ScopingResult(strategy, scoped_allow_list, blocked_tools, ...)",
        "Backend: Dynamic Scoping Engine"
    )

    # 5. Tool Registry Page
    add_custom_heading(doc, "5. Page 3: Tool Registry & Policy Management", level=1)
    add_p(
        doc,
        "The Tool Registry displays all 10 registered capabilities with risk scores (1–10), sensitivity ratings, required roles, and permitted verbs. It includes a direct sandboxed test invocation runner (fulfilling Phase 2 completion evidence) and an admin configuration drawer."
    )
    add_code_block(
        doc,
        "// src/pages/ToolRegistryPage.jsx - Direct Tool Testing Modal\n"
        "const executeDirectTest = async () => {\n"
        "  setTesting(true);\n"
        "  try {\n"
        "    const parsedArgs = JSON.parse(testArgs);\n"
        "    const res = await api.testTool(testModal.name, { arguments: parsedArgs });\n"
        "    setTestResult(res);\n"
        "  } catch (err) {\n"
        "    setTestResult({ error: err.message });\n"
        "  }\n"
        "  setTesting(false);\n"
        "};",
        "Frontend: Tool Registry Direct Invocation"
    )
    add_code_block(
        doc,
        "# backend/app/tool_registry.py - Sandboxed Execution Handlers\n"
        "def handle_calculate_statistics(args: Dict[str, Any]) -> Dict[str, Any]:\n"
        "    numbers = args.get('numbers', [])\n"
        "    arr = np.array([float(x) for x in numbers])\n"
        "    results = { 'count': len(arr), 'mean': float(np.mean(arr)), 'sum': float(np.sum(arr)) }\n"
        "    if args.get('operation') == 'margin' and len(arr) >= 2:\n"
        "        results['margin_percentage'] = round(((arr[0] - arr[1]) / arr[0]) * 100, 2)\n"
        "    return { 'success': True, 'metrics': results }\n\n"
        "def handle_delete_file(args: Dict[str, Any]) -> Dict[str, Any]:\n"
        "    return { 'success': True, 'status': 'DELETED_SANDBOX', 'filepath': args.get('filepath') }",
        "Backend: Sandboxed Tool Implementations"
    )

    # 6. Experiment Lab Page
    add_custom_heading(doc, "6. Page 4: Experiment Lab & Matrix Runner", level=1)
    add_p(
        doc,
        "The Experiment Lab enables researchers to select benchmark tasks, configure the evaluation matrix across strategies (Full, Moderate, Least-Privilege) and model scales (3B vs. 70B), and execute batch tests with live progress. Results populate a comparative findings table."
    )
    add_code_block(
        doc,
        "// src/pages/ExperimentLabPage.jsx - Batch Matrix Execution\n"
        "const handleRunBatch = async () => {\n"
        "  setRunning(true);\n"
        "  setRunProgress('Executing controlled benchmark matrix across configurations...');\n"
        "  const res = await api.runExperiments({\n"
        "    benchmark_ids: selectedBenchmarkIds,\n"
        "    strategies: strategies,\n"
        "    models: models,\n"
        "    repetitions: repetitions\n"
        "  });\n"
        "  setRunProgress(`Completed ${res.total_executions} runs across configurations!`);\n"
        "  const updated = await api.getMetricsSummary();\n"
        "  setMatrixData(updated.comparison_matrix);\n"
        "  setRunning(false);\n"
        "};",
        "Frontend: Experiment Lab Batch Runner"
    )
    add_code_block(
        doc,
        "# backend/app/main.py - Experiment Batch Route\n"
        "@app.post('/api/experiments/run')\n"
        "def run_experiment_batch(req: ExperimentRequest, db: Session = Depends(get_db),\n"
        "                         user: dict = Depends(require_role(UserRole.RESEARCHER))):\n"
        "    benchmarks = db.query(BenchmarkTaskModel).filter(BenchmarkTaskModel.id.in_(req.benchmark_ids)).all()\n"
        "    runs = []\n"
        "    for b in benchmarks:\n"
        "        for strat in req.strategies:\n"
        "            for mod in req.models:\n"
        "                res = AgentRunner.run_task(db, f'EXP-{uuid.uuid4().hex[:8]}', b.task_prompt, strat, mod, ...)\n"
        "                runs.append(res.task_id)\n"
        "    db.commit()\n"
        "    return { 'success': True, 'total_executions': len(runs) }",
        "Backend: Experiment Batch Execution Endpoint"
    )

    # 7. Security Testing Page
    add_custom_heading(doc, "7. Page 5: Controlled Adversarial Security Testing", level=1)
    add_p(
        doc,
        "The Security Testing Suite provides dedicated testing of prompt injection and privilege escalation attack vectors. Clicking 'Run Side-by-Side Adversarial Test' simultaneously dispatches the attack payload through Full Access, Moderate Access, and Least-Privilege, demonstrating how Least-Privilege successfully neutralizes the exploit."
    )
    add_code_block(
        doc,
        "// src/pages/SecurityTestingPage.jsx - 3-Way Comparative Defense Posture Display\n"
        "['Full Access', 'Moderate Access', 'Least-Privilege'].map((strat) => {\n"
        "  const res = comparisonResult[strat];\n"
        "  const attackSucceeded = res.attack_succeeded;\n"
        "  return (\n"
        "    <div className={attackSucceeded ? 'border-red-500 glow-danger' : 'border-emerald-500 glow-success'}>\n"
        "      <span>{strat}</span>\n"
        "      <span className={attackSucceeded ? 'bg-red-600' : 'bg-emerald-600'}>\n"
        "        {attackSucceeded ? 'ATTACK SUCCEEDED (Vulnerable)' : 'ATTACK PREVENTED (Immune)'}\n"
        "      </span>\n"
        "      <p>{attackSucceeded ? '⚠️ Malicious tool executed!' : '🛡️ Scoping stripped tool from allow-list.'}</p>\n"
        "    </div>\n"
        "  );\n"
        "});",
        "Frontend: Side-by-Side Defense Comparison UI"
    )
    add_code_block(
        doc,
        "# backend/app/agent_runner.py - Backend Mediation & Attack Detection\n"
        "if tool_name not in scoped_allow_list:\n"
        "    # Zero-Bypass Boundary: Rejected before LLM call can reach tool\n"
        "    return PolicyDecision.DENY, None, 'ERR_NOT_IN_TASK_SCOPE', 'Tool not in active allow-list.'\n\n"
        "# If tool executes successfully during an adversarial task:\n"
        "if is_adversarial and t_name in ['send_email', 'delete_file', 'modify_database', 'admin_system_command']:\n"
        "    attack_succeeded = True  # Exploited under Full Access!",
        "Backend: Security Mediation Enforcement"
    )

    # 8. Audit Logs Page
    add_custom_heading(doc, "8. Page 6: Tamper-Evident Audit Logs", level=1)
    add_p(
        doc,
        "The Audit Logs view displays a real-time, searchable journal of all tool invocation attempts, accompanied by decision status badges (ALLOW, DENY, REQUIRE_APPROVAL), reason codes, and latency metrics. Clicking 'Details' reveals the sanitized arguments where secrets and passwords have been automatically redacted."
    )
    add_code_block(
        doc,
        "// src/pages/AuditLogsPage.jsx - Audit Log Inspection Drawer\n"
        "{isExpanded && (\n"
        "  <tr className='bg-slate-950'>\n"
        "    <td colSpan={9} className='p-4'>\n"
        "      <span className='font-bold text-slate-400'>Reason: {log.reason_description}</span>\n"
        "      <pre className='text-slate-300 font-mono'>\n"
        "        Sanitized Args: {JSON.stringify(log.sanitized_arguments, null, 2)}\n"
        "      </pre>\n"
        "    </td>\n"
        "  </tr>\n"
        ")}",
        "Frontend: Audit Log Inspection Drawer"
    )
    add_code_block(
        doc,
        "# backend/app/audit_logger.py - Secret-Scrubbing Sanitization Algorithm\n"
        "SENSITIVE_KEYS = {'password', 'secret', 'token', 'auth', 'key', 'api_key', 'db_admin_password'}\n\n"
        "def sanitize_payload(obj: Any) -> Any:\n"
        "    if isinstance(obj, dict):\n"
        "        cleaned = {}\n"
        "        for k, v in obj.items():\n"
        "            if any(s in k.lower() for s in SENSITIVE_KEYS):\n"
        "                cleaned[k] = '[REDACTED_SECRET]'\n"
        "            else:\n"
        "                cleaned[k] = sanitize_payload(v)\n"
        "        return cleaned\n"
        "    return obj",
        "Backend: Automatic Secrets Scrubbing in Audit Trail"
    )

    # 9. Results & Reports Page
    add_custom_heading(doc, "9. Page 7: Research Results & Reports (All 7 Metrics)", level=1)
    add_p(
        doc,
        "The Results & Reports page provides comprehensive visualizations powered by Recharts, evaluating all 7 core research metrics: Task Success Rate, Tool Selection Accuracy, Unnecessary Calls, Security Exposure, Attack Success Rate, Latency, and Token Footprint. Users can download complete CSV and JSON datasets."
    )
    add_code_block(
        doc,
        "// src/pages/ReportsPage.jsx - Recharts Visualization of Security vs Attack Success\n"
        "<ResponsiveContainer width='100%' height={260}>\n"
        "  <BarChart data={securityComparisonData}>\n"
        "    <CartesianGrid strokeDasharray='3 3' stroke='#1e293b' />\n"
        "    <XAxis dataKey='name' stroke='#64748b' />\n"
        "    <YAxis stroke='#64748b' />\n"
        "    <Tooltip contentStyle={{ backgroundColor: '#0f172a' }} />\n"
        "    <Bar dataKey='exposure' name='Security Exposure (0-50)' fill='#f59e0b' />\n"
        "    <Bar dataKey='attackSuccess' name='Attack Success Rate (%)' fill='#ef4444' />\n"
        "  </BarChart>\n"
        "</ResponsiveContainer>",
        "Frontend: Recharts Visualization"
    )
    add_code_block(
        doc,
        "# backend/app/main.py - Exporting Data as CSV or JSON\n"
        "@app.get('/api/reports/export')\n"
        "def export_report(fmt: str = Query('json', enum=['json', 'csv']), db: Session = Depends(get_db)):\n"
        "    summary = EvaluationEngine.calculate_summary(db)\n"
        "    if fmt == 'json':\n"
        "        return summary\n"
        "    output = io.StringIO()\n"
        "    writer = csv.writer(output)\n"
        "    matrix = summary.get('comparison_matrix', [])\n"
        "    if matrix:\n"
        "        writer.writerow(matrix[0].keys())\n"
        "        for row in matrix: writer.writerow(row.values())\n"
        "    return Response(content=output.getvalue(), media_type='text/csv',\n"
        "                    headers={'Content-Disposition': 'attachment; filename=scoping_metrics.csv'})",
        "Backend: CSV & JSON Report Export API"
    )

    # 10. System Settings & RBAC Page
    add_custom_heading(doc, "10. Page 8: System Settings & Role Administration", level=1)
    add_p(
        doc,
        "The Settings page allows Administrators to adjust risk score thresholds via interactive sliders, toggle semantic relevance checks, and enforce exfiltration blockers. It also documents the Role-Based Access Control matrix and runtime environment details."
    )
    add_code_block(
        doc,
        "// src/pages/SettingsPage.jsx - Dynamic Policy Threshold Sliders\n"
        "<div>\n"
        "  <label>Max Risk Score Threshold: {policySettings.max_risk_threshold} / 10</label>\n"
        "  <input\n"
        "    type='range' min='1' max='10'\n"
        "    disabled={!isAdmin}\n"
        "    value={policySettings.max_risk_threshold}\n"
        "    onChange={(e) => setPolicySettings({ ...policySettings, max_risk_threshold: parseInt(e.target.value) })}\n"
        "  />\n"
        "</div>",
        "Frontend: Risk Threshold Adjustment"
    )
    add_code_block(
        doc,
        "# backend/app/policy_engine.py - Policy Threshold Interception\n"
        "if tool.risk_score > max_risk_threshold and user_role != UserRole.ADMIN:\n"
        "    return PolicyEvaluationResult(\n"
        "        decision=PolicyDecision.DENY,\n"
        "        reason_code='ERR_RISK_EXCEEDS_POLICY_THRESHOLD',\n"
        "        reason_description=f'Tool risk ({tool.risk_score}) exceeds safe threshold ({max_risk_threshold}).'\n"
        "    )",
        "Backend: Risk Score Policy Verification"
    )

    # 11. Core Engine Deep-Dive
    add_custom_heading(doc, "11. Core Engine Deep-Dive Source Code", level=1)
    add_p(
        doc,
        "The core defensive intelligence of the framework is implemented across four foundational Python modules:"
    )

    add_code_block(
        doc,
        "# backend/app/task_analyzer.py - Regex Pattern Scanning & Capability Mapping\n"
        "ADVERSARIAL_PATTERNS = [\n"
        "    (r'ignore\\s+(all\\s+)?(previous|prior)\\s+instructions?', 'INJECTION_OVERRIDE_INSTRUCTION'),\n"
        "    (r'(delete|drop|wipe|truncate)\\s+(database|table|users)', 'DESTRUCTIVE_DATABASE_ATTEMPT'),\n"
        "    (r'(send|forward|leak|exfiltrate)\\s+(email|credential|password)\\s+to', 'DATA_EXFILTRATION_ATTEMPT'),\n"
        "    (r'(sudo|root|rm\\s+-rf|cmd\\.exe|powershell)', 'PRIVILEGE_ESCALATION_ATTEMPT')\n"
        "]\n\n"
        "class TaskAnalyzer:\n"
        "    @staticmethod\n"
        "    def analyze(task_prompt: str) -> TaskAnalysisResult:\n"
        "        text = task_prompt.lower()\n"
        "        risk_flags = [flag for pat, flag in ADVERSARIAL_PATTERNS if re.search(pat, text, re.I)]\n"
        "        candidate_tools = [m['tool'] for m in CAPABILITY_MAPPINGS if any(k in text for k in m['keywords'])]\n"
        "        return TaskAnalysisResult(task_prompt, candidate_tools=candidate_tools, risk_flags=risk_flags, ...)",
        "Core Engine: Task Analyzer & Injection Classifier"
    )

    add_code_block(
        doc,
        "# backend/app/main.py - Unified Production SPA Server\n"
        "DIST_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', 'frontend', 'dist'))\n"
        "if os.path.isdir(DIST_DIR):\n"
        "    app.mount('/assets', StaticFiles(directory=os.path.join(DIST_DIR, 'assets')), name='assets')\n"
        "    @app.get('/{full_path:path}')\n"
        "    async def serve_spa(full_path: str):\n"
        "        if full_path.startswith('api'): raise HTTPException(status_code=404)\n"
        "        file_path = os.path.join(DIST_DIR, full_path)\n"
        "        if os.path.isfile(file_path): return FileResponse(file_path)\n"
        "        return FileResponse(os.path.join(DIST_DIR, 'index.html'))",
        "Core Engine: Single-Port Production Serving"
    )

    doc.save(DOC2_PATH)
    print(f"Successfully generated: {DOC2_PATH}")


if __name__ == "__main__":
    build_document_2()
